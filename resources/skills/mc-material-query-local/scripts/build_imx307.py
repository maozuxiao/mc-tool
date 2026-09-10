# -*- coding: utf-8 -*-
"""
从 IMX307替代料号明细.xlsx 重新生成 data/imx307_replacement.json

用法:
  python build_imx307.py                 # 自动选取 AHD 目录下最新的 IMX307*.xlsx
  python build_imx307.py <xlsx路径>       # 指定源文件

数据源列映射（0 起）:
  翔飞(8列): 编号->original, 描述->original_desc, 生命周期阶段->lifecycle,
             机型->model, 客户代码->customer_code, 替代料号->replacement, 描述->replacement_desc
  罡扇(5列): 编号->original, 描述->original_desc, 生命周期阶段->lifecycle,
             替代料号->replacement, 描述->replacement_desc  (无 机型/客户代码 列)

规则:
  - 跳过「编号」或「替代料号」为空的行（保持"命中映射=有替代"的既有语义，
    未分配替代料号的 IMX307 料号不进入映射表，查询时表现为"已检查但无替代记录"）
  - 每条目新增 supplier 字段（翔飞=F355方案 / 罡扇=2053方案）区分替代来源供应商
"""
import os, sys, glob, json

DEFAULT_DIR = r"E:\毛祖潇的知识库\产品资料\产品线\摄像机\AHD"
OUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "imx307_replacement.json")

# 输出字段顺序（前 7 项与历史 JSON 保持一致，supplier 为本次新增）
FIELDS = ["original", "original_desc", "replacement", "replacement_desc",
          "lifecycle", "model", "customer_code"]

COL_IDX = {
    "翔飞": {"original": 0, "original_desc": 1, "lifecycle": 2, "model": 3,
             "customer_code": 4, "replacement": 5, "replacement_desc": 6},
    "罡扇": {"original": 0, "original_desc": 1, "lifecycle": 2,
             "replacement": 3, "replacement_desc": 4},
}


def cell(v):
    """单元格值 -> 规范字符串或 None。料号保持文本、数字去掉 .0，避免 13 位料号被转成科学计数/丢尾。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s if s != "" else None


def pick_src(argv):
    if len(argv) > 1:
        return argv[1]
    files = [f for f in glob.glob(os.path.join(DEFAULT_DIR, "IMX307*.xlsx"))
             if not os.path.basename(f).startswith("~$")]
    if not files:
        raise SystemExit("未在 %s 找到 IMX307*.xlsx" % DEFAULT_DIR)
    files.sort(key=lambda f: os.path.getmtime(f), reverse=True)
    return files[0]


def main():
    src = pick_src(sys.argv)
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True)

    mapping = []
    for sheet, idx in COL_IDX.items():
        if sheet not in wb.sheetnames:
            print("[WARN] 缺少工作表:", sheet)
            continue
        ws = wb[sheet]
        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            original = cell(row[idx["original"]]) if idx["original"] < len(row) else None
            replacement = cell(row[idx["replacement"]]) if idx["replacement"] < len(row) else None
            if not original or not replacement:
                continue
            entry = {}
            for f in FIELDS:
                entry[f] = cell(row[idx[f]]) if f in idx and idx[f] < len(row) else None
            entry["supplier"] = sheet
            mapping.append(entry)
            n += 1
        print("[OK] %s: %d 条" % (sheet, n))

    out = {
        "source": src,
        "note": "IMX307方案料号 -> 替代料号 映射表（静态参考，生命周期/库存以MC服务器为准）。翔飞=F355方案，罡扇=2053方案。",
        "count": len(mapping),
        "mapping": mapping,
    }
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print("[DONE] 共 %d 条 -> %s" % (len(mapping), OUT_PATH))


if __name__ == "__main__":
    main()
